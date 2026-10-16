"""Import groups from pasted TSV/CSV text or an .xlsx upload.

Pure functions, no database access. Output times are minutes since midnight
and dates are ISO strings (see ``timeutil``).
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from dataclasses import dataclass

from timeutil import fmt_time, parse_date, parse_time

HEADER_ALIASES = {
    "name": ("group", "name", "activity", "title", "booking"),
    "date": ("date", "day"),
    "start": ("start", "from", "start time", "begin"),
    "end": ("end", "to", "end time", "until", "finish"),
    "duration": ("duration", "length", "hours", "duration (h)"),
    "size": ("size", "people", "participants", "headcount", "guests"),
}

HEADER_ERROR = ("Couldn't find columns for date, start and end. "
                "Use a header row like: Group, Date, Start, End")

_SECONDS_RE = re.compile(r"^(\d{1,2}[:.]\d{2}):\d{2}$")
_DUR_HM_RE = re.compile(r"^(\d{1,2}):(\d{2})$")
_DUR_MIN_RE = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(m|min|mins|minutes?)$")
_DUR_H_RE = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(h|hr|hrs|hours?)?$")


@dataclass
class GroupRow:
    name: str
    date: str
    start: int
    end: int
    size: int | None
    line: int


@dataclass
class RowError:
    line: int
    message: str


# --------------------------------------------------------------------------- #
# Cell parsing helpers
# --------------------------------------------------------------------------- #

def _parse_time_cell(s: str) -> int:
    text = str(s).strip()
    m = _SECONDS_RE.match(text)
    if m:  # "10:00:00" as pasted from Excel
        text = m.group(1)
    return parse_time(text)


def _parse_duration(s: str) -> int:
    """Duration -> minutes. Accepts "1.5", "1,5", "1:30", "90 min", "2h"."""
    text = str(s).strip().lower()
    m = _DUR_HM_RE.match(text)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    m = _DUR_MIN_RE.match(text)
    if m:
        return int(round(float(m.group(1).replace(",", "."))))
    m = _DUR_H_RE.match(text)
    if m:
        return int(round(float(m.group(1).replace(",", ".")) * 60))
    raise ValueError(f"Duration '{s}' should look like 1:30 or 1.5")


def _parse_size(s: str) -> int | None:
    text = str(s).strip()
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        raise ValueError(f"Size '{s}' should be a whole number") from None


def _map_header(cells: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(cells):
        key = str(cell).strip().lower()
        if not key:
            continue
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases and field_name not in mapping:
                mapping[field_name] = idx
                break
    if "date" in mapping and "start" in mapping and ("end" in mapping or "duration" in mapping):
        return mapping
    return None


def _cell(row: list[str], mapping: dict[str, int], field_name: str) -> str:
    idx = mapping.get(field_name)
    if idx is None or idx >= len(row):
        return ""
    return str(row[idx]).strip()


# --------------------------------------------------------------------------- #
# Core conversion
# --------------------------------------------------------------------------- #

def rows_to_groups(rows: list[list[str]]) -> tuple[list[GroupRow], list[RowError]]:
    """Convert raw string rows (header first) into GroupRows plus per-row errors."""
    groups: list[GroupRow] = []
    errors: list[RowError] = []

    header_idx = None
    for i, row in enumerate(rows):
        if any(str(c).strip() for c in row):
            header_idx = i
            break
    if header_idx is None:
        return [], [RowError(1, HEADER_ERROR)]

    mapping = _map_header([str(c) for c in rows[header_idx]])
    if mapping is None:
        return [], [RowError(1, HEADER_ERROR)]

    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        line = i + 1
        if not any(str(c).strip() for c in row):
            continue
        row_errors: list[str] = []

        name = _cell(row, mapping, "name") or "Group"

        date = None
        try:
            date = parse_date(_cell(row, mapping, "date"))
        except ValueError as e:
            row_errors.append(str(e))

        start = None
        try:
            start = _parse_time_cell(_cell(row, mapping, "start"))
        except ValueError as e:
            row_errors.append(str(e))

        end = None
        end_text = _cell(row, mapping, "end")
        dur_text = _cell(row, mapping, "duration")
        if end_text:
            try:
                end = _parse_time_cell(end_text)
            except ValueError as e:
                row_errors.append(str(e))
        elif dur_text:
            try:
                minutes = _parse_duration(dur_text)
                if start is not None:
                    end = start + minutes
            except ValueError as e:
                row_errors.append(str(e))
        else:
            row_errors.append("End time or duration is missing")

        size = None
        try:
            size = _parse_size(_cell(row, mapping, "size"))
        except ValueError as e:
            row_errors.append(str(e))

        if start is not None and end is not None and end <= start:
            row_errors.append(
                f"end time {fmt_time(end % 1440)} is not after start time {fmt_time(start)}")

        if row_errors:
            for msg in row_errors:
                errors.append(RowError(line, f"Row {line}: {msg}"))
            continue

        groups.append(GroupRow(name, date, start, end, size, line))

    return groups, errors


# --------------------------------------------------------------------------- #
# Entry points
# --------------------------------------------------------------------------- #

def _detect_delimiter(text: str) -> str:
    for line in text.splitlines():
        if line.strip():
            if "\t" in line:
                return "\t"
            return ";" if line.count(";") > line.count(",") else ","
    return ","


def parse_text(text: str) -> tuple[list[GroupRow], list[RowError]]:
    """Parse pasted TSV/CSV text."""
    text = (text or "").lstrip("\ufeff")
    delimiter = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = [[c for c in row] for row in reader]
    return rows_to_groups(rows)


def _xlsx_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        # Excel stores time-only cells as datetimes on the epoch day.
        if value.year < 1900:
            return fmt_time(value.hour * 60 + value.minute)
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, _dt.time):
        return fmt_time(value.hour * 60 + value.minute)
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if 0 <= value < 1:  # Excel time fraction of a day
            return fmt_time(int(round(value * 1440)))
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def parse_xlsx(data: bytes) -> tuple[list[GroupRow], list[RowError]]:
    """Parse the first sheet of an .xlsx workbook."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        return [], [RowError(1, "This file doesn't look like an Excel (.xlsx) workbook.")]
    try:
        ws = wb.worksheets[0]
        rows = [[_xlsx_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows_to_groups(rows)
